import os, json, sqlite3, time, uuid, hashlib, secrets, re, io, base64, zipfile, threading
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlencode
import requests
from fastapi import FastAPI, Request, UploadFile, File, Form, HTTPException
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse
from cryptography.fernet import Fernet
from docx import Document
from pypdf import PdfReader
from starlette.concurrency import run_in_threadpool
from mcp_connections import Connections, ConnectionError, WriteUncertainError
from mcp_services import Services, notion_id, FOLDER, scalar, properties
from dotenv import load_dotenv

ROOT=Path(__file__).parent
load_dotenv(ROOT/'.env')
DATA=Path(os.getenv('DATA_DIR',str(ROOT/'local-data'))); DATA.mkdir(parents=True,exist_ok=True)
DB=DATA/'portal.sqlite3'
UPLOADS=DATA/'uploads'; UPLOADS.mkdir(exist_ok=True)
def config(k,default=''): return os.getenv(k,default)
def now(): return datetime.now(timezone.utc).isoformat()
def conn():
 c=sqlite3.connect(DB,timeout=30); c.row_factory=sqlite3.Row; return c
def execute(q,params=()):
 with conn() as c: return c.execute(q,params)
def rows(q,params=()):
 with conn() as c: return [dict(x) for x in c.execute(q,params)]
def one(q,params=()):
 with conn() as c:
  x=c.execute(q,params).fetchone(); return dict(x) if x else None
with conn() as c:
 c.execute('PRAGMA journal_mode=WAL')
 c.executescript('''CREATE TABLE IF NOT EXISTS partners(id TEXT PRIMARY KEY,name TEXT NOT NULL,queue TEXT NOT NULL DEFAULT 'To pick up',stage TEXT NOT NULL DEFAULT 'Not received',notes TEXT NOT NULL DEFAULT '',owner TEXT DEFAULT '',notion_id TEXT,notion_url TEXT,contract_url TEXT,verification INTEGER DEFAULT 0,sync_state TEXT DEFAULT 'demo',updated TEXT);
 CREATE TABLE IF NOT EXISTS files(id TEXT PRIMARY KEY,partner_id TEXT,name TEXT,hash TEXT,drive_id TEXT,drive_url TEXT,path TEXT,document_type TEXT,source TEXT DEFAULT 'portal',sync_state TEXT,created TEXT,analysis TEXT);
 CREATE TABLE IF NOT EXISTS deviations(id TEXT PRIMARY KEY,file_id TEXT,subject TEXT,base_text TEXT,partner_text TEXT,decision TEXT DEFAULT 'Needs review',explanation TEXT,reviewer TEXT DEFAULT '',reviewed TEXT);
 CREATE TABLE IF NOT EXISTS events(id TEXT PRIMARY KEY,partner_id TEXT,actor TEXT,action TEXT,detail TEXT,created TEXT);
 CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY,value TEXT);
 CREATE INDEX IF NOT EXISTS ix_partner_name ON partners(name); CREATE INDEX IF NOT EXISTS ix_file_hash ON files(partner_id,hash);''')
 columns={row[1] for row in c.execute('PRAGMA table_info(partners)')}
 if 'source' not in columns:
  c.execute("ALTER TABLE partners ADD COLUMN source TEXT DEFAULT 'local'")
  c.execute("UPDATE partners SET source='demo' WHERE sync_state='demo'")
 if 'pending_properties' not in columns:c.execute("ALTER TABLE partners ADD COLUMN pending_properties TEXT DEFAULT '{}'")
 if 'notion_source_id' not in columns:
  c.execute("ALTER TABLE partners ADD COLUMN notion_source_id TEXT DEFAULT ''")
  c.execute("UPDATE partners SET notion_source_id=? WHERE notion_id IS NOT NULL",(config('NOTION_DATA_SOURCE_ID'),))
  c.execute("UPDATE partners SET source='notion' WHERE notion_id IS NOT NULL AND source!='demo'")
 file_columns={row[1] for row in c.execute('PRAGMA table_info(files)')}
 if 'note' not in file_columns:c.execute("ALTER TABLE files ADD COLUMN note TEXT DEFAULT ''")
 if not c.execute('SELECT 1 FROM partners LIMIT 1').fetchone() and config('DEMO_MODE','0')=='1':
  seed=[('Demo Partner A','To pick up','Not received','Example record'),('Demo Partner B','Saket review','Negotiating','Example decision pending'),('Demo Partner C','On hold','With partner','Example paused partnership')]
  for name,queue,stage,notes in seed: c.execute('INSERT INTO partners(id,name,queue,stage,notes,updated,verification) VALUES(?,?,?,?,?,?,?)',(str(uuid.uuid4()),name,queue,stage,notes,now(),0))
  c.execute("UPDATE partners SET source='demo' WHERE sync_state='demo'")

def setting(key, default=''):
 x=one('SELECT value FROM settings WHERE key=?',(key,)); return x['value'] if x else default
def put_setting(k,v): execute('INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value',(k,v))
def log(pid,action,detail,actor='Finance'):
 execute('INSERT INTO events VALUES(?,?,?,?,?,?)',(str(uuid.uuid4()),pid,actor,action,detail,now()))

connections=Connections(DB,DATA,config)
services=Services(connections,setting,put_setting,config,rows,one,execute,log)
operation_lock=threading.RLock()

def notion_sync():return services.sync_notion()
def notion_update(pid,properties):return services.update_notion(pid,properties)
def google_folder_scan():return services.sync_drive()
def upload_drive(filename,raw,mime):return services.upload_drive(filename,raw,mime)
def google_file_bytes(fid):return services.file_bytes(fid)
def token():return connections.connected('google')

def connection_summary():
 n=connections.status('notion');g=connections.status('google')
 return {'notion':n['connected'],'google':g['connected'],'notion_verified':n['verified'],'google_verified':g['verified'],'notion_status':n,'google_status':g,'notion_last_sync':setting('notion_sync_at'),'drive_last_sync':setting('drive_sync_at')}

def record_pending(pid,values):
 p=one('SELECT pending_properties FROM partners WHERE id=?',(pid,))
 pending=json.loads(p['pending_properties'] or '{}');pending.update(values)
 execute('UPDATE partners SET pending_properties=? WHERE id=?',(json.dumps(pending),pid))

def extract(raw,filename):
 if filename.lower().endswith('.pdf'):
  rd=PdfReader(io.BytesIO(raw));parts=[]
  for i,p in enumerate(rd.pages):
   tx=p.extract_text() or ''
   if tx.strip():parts.append((i+1,tx))
  if not parts:raise RuntimeError('No machine-readable text found. OCR/manual review required.')
  return parts
 if filename.lower().endswith('.docx'):
  d=Document(io.BytesIO(raw));paras=[x.text for x in d.paragraphs if x.text.strip()]
  for tbl in d.tables:
   for r in tbl.rows:
    paras.append(' | '.join(c.text for c in r.cells))
  if not paras:raise RuntimeError('DOCX is empty or unreadable')
  return [(1,'\n'.join(paras))]
 raise RuntimeError('Unsupported file; PDF and DOCX only')
BASE_RULES=[('Commission rates',r'commission|revenue share|reseller discount|margin','Compare commission tiers and revenue definitions against the approved template.'),('Post-termination tail',r'post.termination|following termination|upon termination|commission payable','Compare post-termination commission protection against the approved template.'),('Termination for convenience',r'termination for convenience|days.notice','Compare notice periods and termination rights against the approved template.'),('Territory',r'territory|jurisdiction','Compare the covered territory against the approved template.'),('Customer payments',r'non.payment|unable to recover|payment terms|payments from customer','Compare payment timing and non-payment protection against the approved template.'),('Product change notice',r'discontinuation|product change|notice of changes','Compare product-change notice requirements against the approved template.'),('Opportunity registration',r'approve or reject|opportunity registration|engagement request','Compare opportunity registration rights against the approved template.'),('Renewals and upsells',r'renewals|upsells|extensions','Compare renewal and expansion rights against the approved template.'),('Indemnification',r'indemni','Compare scope, exclusions, and caps against the approved template.'),('Liability',r'limitation of liability|liability cap','Compare caps, exclusions, and mutuality against the approved template.')]

def snippets(parts,regex,limit=1100):
 patt=re.compile(regex,re.I);found=[]
 for page,tx in parts:
  for m in list(patt.finditer(tx))[:3]:found.append(f'Page {page}: '+re.sub(r'\s+',' ',tx[max(0,m.start()-220):m.end()+max(250,limit//3)]).strip()[:limit]);
  if len(found)>=3:break
 return '\n\n'.join(found[:2])
def compare(file_id,raw,name,doctype):
 parts=extract(raw,name)
 template_id=config('RESELLER_TEMPLATE_ID') if doctype!='LOI' else config('LOI_TEMPLATE_ID')
 base=None; basis='manual review checklist (approved template unavailable)'
 if template_id and token():
  try:base=extract(google_file_bytes(template_id), 'base.docx');basis='live Drive template'
  except Exception:pass
 execute('DELETE FROM deviations WHERE file_id=?',(file_id,))
 total=0
 for title,pattern,summary in BASE_RULES:
  partner_text=snippets(parts,pattern)
  if not partner_text:continue
  base_text=snippets(base,pattern) if base else summary
  # No legal implication inferred by fuzzy matching; reviewers get both exact excerpts.
  if base and re.sub(r'\W+','',partner_text.lower())==re.sub(r'\W+','',base_text.lower()):decision='Aligned'
  else:decision='Needs review'
  extra='Actual source extracts shown. Semantic differences require human assessment. '
  execute('INSERT INTO deviations(id,file_id,subject,base_text,partner_text,decision,explanation) VALUES(?,?,?,?,?,?,?)',(str(uuid.uuid4()),file_id,title,base_text[:3500],partner_text[:3500],decision,extra+'Baseline: '+basis));total+=1
 return {'sections_extracted':len(parts),'review_items':total,'basis':basis,'disclaimer':'Heuristic extraction, not legal review'}

if config('APP_ENV')=='production' and not config('PORTAL_PASSWORD'):raise RuntimeError('PORTAL_PASSWORD required in production')
app=FastAPI(title='Onnivation Partnership Command Center')
@app.middleware('http')
async def auth(request:Request,call_next):
 if request.url.path in ('/health','/favicon.ico','/auth/google/callback','/auth/notion/callback'):return await call_next(request)
 password=config('PORTAL_PASSWORD')
 if not password:
  # Local-only startup still requires explicit secrets for any external access.
  if request.client and request.client.host not in ('127.0.0.1','::1','localhost'):return JSONResponse({'error':'Set PORTAL_PASSWORD before exposing server'},status_code=503)
 else:
  auth=request.headers.get('Authorization','')
  try:
   typ,encoded=auth.split(' ',1);username,pwd=base64.b64decode(encoded).decode().split(':',1)
   valid=typ.lower()=='basic' and secrets.compare_digest(username,config('PORTAL_USERNAME','finance')) and secrets.compare_digest(pwd,password)
  except Exception:valid=False
  if not valid:return JSONResponse({'error':'Authentication required'},status_code=401,headers={'WWW-Authenticate':'Basic realm="Onnivation Finance Portal"'})
 if request.method in ('POST','PATCH','DELETE') and request.url.path not in ('/auth/google/callback',):
  # Cross-origin writes blocked in addition to basic auth, preventing simple CSRF.
  origin=request.headers.get('origin')
  if origin and origin.rstrip('/')!=connections.base_url():return JSONResponse({'error':'Origin rejected'},status_code=403)
  if request.headers.get('sec-fetch-site')=='cross-site':return JSONResponse({'error':'Cross-site request rejected'},status_code=403)
 response=await call_next(request)
 response.headers['Cache-Control']='no-store'
 response.headers['X-Content-Type-Options']='nosniff'
 response.headers['Referrer-Policy']='no-referrer'
 response.headers['X-Frame-Options']='DENY'
 return response
@app.get('/')
def index():return FileResponse(ROOT/'static'/'index.html')
@app.get('/health')
def health():return {'ok':True}
@app.get('/api/overview')
def overview():
 summary=connection_summary()
 real=bool(one("SELECT id FROM partners WHERE source!='demo' LIMIT 1")) or summary['notion']
 condition=" WHERE source!='demo' AND (source='local' OR notion_source_id=?)" if real else ''
 p=rows('SELECT * FROM partners'+condition+' ORDER BY CASE queue WHEN "Saket review" THEN 0 WHEN "To pick up" THEN 1 ELSE 2 END, name',(services.source(),) if real else ())
 f=rows('SELECT id,partner_id,name,drive_url,source,sync_state,created,document_type,note FROM files ORDER BY created DESC LIMIT 200')
 events=rows('SELECT e.*,p.name as partner FROM events e LEFT JOIN partners p ON p.id=e.partner_id ORDER BY e.created DESC LIMIT 100')
 mode='live' if summary['notion_verified'] and summary['google_verified'] else 'partial' if summary['notion'] or summary['google'] else 'local' if real or not p else 'demo'
 return {'partners':p,'files':f,'events':events,'deviation_count':one('SELECT COUNT(*) as n FROM deviations WHERE decision="Needs review"')['n'],'integrations':summary,'mode':mode,'settings':{'notion_data_source_id':services.source(),'drive_folder_id':services.folder(),'upload_folder_id':services.folder(True)}}

def change_partner(pid,field,value,reason=''):
 with operation_lock:
  p=one('SELECT * FROM partners WHERE id=?',(pid,))
  if not p:raise HTTPException(404,'Partner not found')
  if p['notion_id'] and p['notion_source_id']!=services.source():raise HTTPException(409,'This partner belongs to another Notion source. Refresh and select a partner from the current source.')
  if p['sync_state']=='publishing':raise HTTPException(409,'Create completion is uncertain. Inspect Notion before editing this draft.')
  prop='Queue' if field=='queue' else 'Contract stage'
  state='demo' if p['source']=='demo' else 'pending';issue=''
  if p['notion_id'] and p['source']!='demo':
   try:notion_update(p['notion_id'],{prop:value});state='live'
   except Exception as e:issue=str(e)
  if state=='pending':record_pending(pid,{prop:value})
  if state=='live':
   pending=json.loads(p['pending_properties'] or '{}');pending.pop(prop,None)
   execute('UPDATE partners SET pending_properties=? WHERE id=?',(json.dumps(pending),pid))
   if pending:state='pending'
  execute('UPDATE partners SET '+field+'=?,sync_state=?,updated=? WHERE id=?',(value,state,now(),pid))
  log(pid,'Queue change' if field=='queue' else 'Contract stage',str(p[field])+' → '+value+' | '+reason[:500]+' | '+state)
  return {field:value,'sync_state':state,'warning':issue or ('Demo change only; never published' if state=='demo' else 'Local draft awaiting Notion' if state=='pending' else '')}

@app.post('/api/partners/{pid}/move')
async def move(pid:str,request:Request):
 payload=await request.json();queue=payload.get('queue')
 if queue not in ('To pick up','Saket review','On hold'):raise HTTPException(400,'Invalid queue')
 return await run_in_threadpool(change_partner,pid,'queue',queue,str(payload.get('reason','')))

@app.post('/api/partners/{pid}/stage')
async def stage(pid:str,request:Request):
 d=await request.json();value=d.get('stage')
 if value not in ('Not received','NDA / LOI','With partner','Negotiating','For signature','Signed'):raise HTTPException(400,'Invalid stage')
 return await run_in_threadpool(change_partner,pid,'stage',value)

@app.post('/api/partners')
async def create_partner(request:Request):
 d=await request.json();name=(d.get('name') or '').strip()
 if not name or len(name)>200:raise HTTPException(400,'Partner name must be 1–200 characters')
 def perform():
  with operation_lock:
   old=one("SELECT * FROM partners WHERE LOWER(name)=LOWER(?) AND source!='demo' AND (source='local' OR notion_source_id=?)",(name,services.source()))
   if old:return {'partner':old,'existing':True}
   pid=str(uuid.uuid4());execute("INSERT INTO partners(id,name,queue,stage,notes,sync_state,updated,source) VALUES(?,?,?,?,?,?,?,?)",(pid,name,'To pick up','Not received','Created from portal','pending',now(),'local'))
   log(pid,'Created','New partner draft')
   warning=''
   if connections.status('notion')['verified']:
    try:services.create_notion(one('SELECT * FROM partners WHERE id=?',(pid,)))
    except Exception as e:warning=str(e)
   return {'partner':one('SELECT * FROM partners WHERE id=?',(pid,)),'existing':False,'warning':warning}
 return await run_in_threadpool(perform)

@app.post('/api/partners/{pid}/resolve')
async def resolve_partner(pid:str,request:Request):
 d=await request.json()
 try:page_id=notion_id(d.get('notion_id',''))
 except ConnectionError as e:raise HTTPException(400,str(e))
 def perform():
  with operation_lock:
   p=one('SELECT * FROM partners WHERE id=?',(pid,))
   if not p or p['sync_state']!='publishing':raise HTTPException(400,'This draft does not need create reconciliation')
   matches=[r for r in services.notion_rows() if notion_id(r.get('id') or r.get('url') or r.get('page_id') or r.get('page_url'))==page_id]
   if len(matches)!=1:raise HTTPException(400,'The page is not in the selected partnership source')
   actual=properties(connections.call('notion','notion-fetch',{'id':page_id}))
   if actual.get('Partner')!=p['name']:raise HTTPException(400,'Notion partner name does not match this draft')
   with conn() as db:
    duplicates=[r['id'] for r in db.execute('SELECT id FROM partners WHERE notion_id=? AND id!=?',(page_id,pid))]
    for duplicate in duplicates:
     db.execute('UPDATE files SET partner_id=? WHERE partner_id=?',(pid,duplicate))
     db.execute('UPDATE events SET partner_id=? WHERE partner_id=?',(pid,duplicate))
     db.execute('DELETE FROM partners WHERE id=?',(duplicate,))
    db.execute("UPDATE partners SET notion_id=?,notion_url=?,notion_source_id=?,source='notion',sync_state='live',pending_properties='{}',queue=?,stage=? WHERE id=?",(page_id,'https://www.notion.so/'+page_id.replace('-',''),services.source(),actual['Queue'],actual['Contract stage'],pid))
   log(pid,'Partner reconciled','Existing Notion page verified; no duplicate created')
   return {'saved':True}
 return await run_in_threadpool(perform)

@app.post('/api/sync')
def sync():
 # Refresh is read-only: pending drafts are published by a separate explicit action.
 with operation_lock:
  out={}
  for label,provider,fn in [('notion','notion',notion_sync),('drive','google',google_folder_scan)]:
   try:out[label]={'success':True,'processed':fn()}
   except Exception as e:
    connections.write(provider+':health',{'verified':False,'error':str(e)[:500],'checked_at':now()})
    out[label]={'success':False,'error':str(e)}
  return out

@app.post('/api/upload')
async def upload(file:UploadFile=File(...),partner_id:str=Form(...),document_type:str=Form('Reseller agreement'),note:str=Form('')):
 p=one('SELECT * FROM partners WHERE id=?',(partner_id,))
 if not p:raise HTTPException(404,'Partner not found')
 if p['notion_id'] and p['notion_source_id']!=services.source():raise HTTPException(409,'Partner belongs to another Notion source')
 if p['source']=='demo':raise HTTPException(400,'Create or import a real partner before uploading contracts')
 if document_type not in ('Reseller agreement','LOI','NDA','Other agreement'):raise HTTPException(400,'Unknown document type')
 fname=Path(file.filename or '').name
 if not fname.lower().endswith(('.pdf','.docx')) or len(fname)>240:raise HTTPException(400,'Choose a PDF or DOCX with a valid filename')
 raw=await file.read(20*1024*1024+1)
 if len(raw)>20*1024*1024:raise HTTPException(413,'20 MB upload limit')
 if fname.lower().endswith('.pdf') and not raw.startswith(b'%PDF'):raise HTTPException(400,'File is not a PDF')
 if fname.lower().endswith('.docx'):
  try:
   with zipfile.ZipFile(io.BytesIO(raw)) as archive:
    if 'word/document.xml' not in archive.namelist():raise ValueError()
    if sum(i.file_size for i in archive.infolist())>80*1024*1024:raise ValueError()
  except Exception:raise HTTPException(400,'Invalid or oversized DOCX contents')
 def perform():
  with operation_lock:
   digest=hashlib.sha256(raw).hexdigest();match=one('SELECT * FROM files WHERE partner_id=? AND hash=?',(partner_id,digest))
   if match:return {'duplicate':True,'file':match,'message':'Identical document is already recorded'}
   fid=str(uuid.uuid4());dest=UPLOADS/(fid+Path(fname).suffix.lower());dest.write_bytes(raw)
   execute('INSERT INTO files(id,partner_id,name,hash,drive_id,drive_url,path,document_type,source,sync_state,created,analysis,note) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)',(fid,partner_id,fname,digest,'','',str(dest),document_type,'portal','local_only',now(),'',note[:2000]))
   result=retry_file(fid)
   try:
    analysis=compare(fid,raw,fname,'LOI' if document_type=='LOI' else document_type) if document_type in ('LOI','Reseller agreement') else {'manual_review':True,'basis':'No approved baseline configured for this document type'}
   except Exception as e:analysis={'error':str(e),'manual_review':True}
   execute('UPDATE files SET analysis=? WHERE id=?',(json.dumps(analysis),fid));log(partner_id,'Uploaded',fname+' | '+result['file']['sync_state']+' | '+note[:500])
   return {'file':one('SELECT * FROM files WHERE id=?',(fid,)),'analysis':analysis,'warnings':result['warnings'],'status':result['file']['sync_state']}
 return await run_in_threadpool(perform)

@app.get('/api/files/{fid}/review')
def review(fid:str):
 f=one('SELECT * FROM files WHERE id=?',(fid,));
 if not f:raise HTTPException(404,'No file')
 return {'file':f,'deviations':rows('SELECT * FROM deviations WHERE file_id=? ORDER BY subject',(fid,))}
@app.post('/api/deviations/{did}/decision')
async def decision(did:str,request:Request):
 d=await request.json();choice=d.get('decision')
 if choice not in ('Approved exception','Needs review','Negotiation requested','Aligned'):raise HTTPException(400,'Invalid review state')
 r=one('SELECT f.partner_id FROM deviations d JOIN files f ON d.file_id=f.id WHERE d.id=?',(did,))
 if not r:raise HTTPException(404,'Review item missing')
 execute('UPDATE deviations SET decision=?,reviewer=?,reviewed=? WHERE id=?',(choice,d.get('reviewer') or 'Finance',now(),did));log(r['partner_id'],'Clause decision',choice+' | '+d.get('reason',''))
 return {'saved':True,'decision':choice}
def retry_file(fid):
 with operation_lock:
  f=one('SELECT * FROM files WHERE id=?',(fid,))
  if not f or not f['path']:raise HTTPException(404,'Local upload not available')
  if f['sync_state']=='synced':return {'success':True,'warnings':[],'file':f}
  if f['sync_state'] in ('drive_uploading','drive_uncertain') and not f['drive_id']:
   return {'success':False,'warnings':['Upload completion is uncertain. Inspect Drive and reconcile the existing file before retrying.'],'file':f}
  p=one('SELECT * FROM partners WHERE id=?',(f['partner_id'],));raw=Path(f['path']).read_bytes();driveid=f['drive_id'];url=f['drive_url'];errors=[]
  if not driveid:
   if not connections.connected('google'):return {'success':False,'warnings':['Connect Google Drive to upload this staged file.'],'file':f}
   try:
    services.verify_folder(True)
    execute("UPDATE files SET sync_state='drive_uploading' WHERE id=?",(fid,))
    res=upload_drive(f['name'],raw,'application/pdf' if f['name'].lower().endswith('.pdf') else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    driveid=res['id'];url=res['webViewLink'];execute('UPDATE files SET drive_id=?,drive_url=?,sync_state=? WHERE id=?',(driveid,url,'drive_uploaded',fid))
   except Exception as e:
    if one('SELECT sync_state FROM files WHERE id=?',(fid,))['sync_state']=='drive_uploading':execute('UPDATE files SET sync_state=? WHERE id=?',('drive_uncertain' if isinstance(e,WriteUncertainError) else 'drive_error',fid))
    errors.append(str(e))
  if url:
   try:
    notion_update(p['notion_id'],{'Contract link':url})
    execute("UPDATE files SET sync_state='synced' WHERE id=?",(fid,))
    execute('UPDATE partners SET contract_url=? WHERE id=?',(url,p['id']))
   except Exception as e:errors.append(str(e));execute("UPDATE files SET sync_state='notion_pending' WHERE id=?",(fid,))
  return {'success':not errors,'warnings':errors,'file':one('SELECT * FROM files WHERE id=?',(fid,))}

@app.post('/api/files/{fid}/retry')
def retry(fid:str):return retry_file(fid)

@app.post('/api/files/{fid}/resolve')
async def resolve_upload(fid:str,request:Request):
 d=await request.json();driveid=str(d.get('drive_id','')).strip()
 if not FOLDER.fullmatch(driveid):raise HTTPException(400,'Enter the existing Drive file ID')
 def perform():
  with operation_lock:
   f=one('SELECT * FROM files WHERE id=?',(fid,))
   if not f or f['sync_state'] not in ('drive_uploading','drive_uncertain'):raise HTTPException(400,'This file does not need upload reconciliation')
   metadata=connections.call('google','get_file_metadata',{'fileId':driveid})
   if metadata.get('parentId')!=services.folder(True):raise HTTPException(400,'File is not in the selected upload folder')
   content=google_file_bytes(driveid)
   if hashlib.sha256(content).hexdigest()!=f['hash']:raise HTTPException(400,'Drive content differs from the uploaded document')
   url=metadata.get('viewUrl') or 'https://drive.google.com/file/d/'+driveid+'/view'
   execute('UPDATE files SET drive_id=?,drive_url=?,sync_state=? WHERE id=?',(driveid,url,'drive_uploaded',fid))
   log(f['partner_id'],'Upload reconciled','Existing Drive file content verified; no duplicate uploaded')
   return retry_file(fid)
 return await run_in_threadpool(perform)

@app.get('/api/settings')
def settings():
 return {**connection_summary(),'notion_data_source_id':services.source(),'drive_folder_id':services.folder(),'upload_folder_id':services.folder(True),'notion_db':setting('notion_database_url',config('NOTION_DATABASE_URL'))}

@app.post('/api/settings')
async def save_settings(request:Request):
 d=await request.json()
 allowed={'notion_data_source_id','drive_folder_id','upload_folder_id'}
 if not set(d).issubset(allowed):raise HTTPException(400,'Unknown setting')
 normalized={}
 for key,value in d.items():
  value=str(value).strip()
  if key=='notion_data_source_id':
   try:value=notion_id(value)
   except ConnectionError as e:raise HTTPException(400,str(e))
  else:
   match=re.search(r'/folders/([A-Za-z0-9_-]+)',value)
   if match:value=match[1]
   if not FOLDER.fullmatch(value):raise HTTPException(400,'Enter a valid Drive folder URL or ID')
  normalized[key]=value
 def perform():
  with operation_lock:
   for key,value in normalized.items():
    if setting(key)!=value:
     put_setting(key,value)
     provider='notion' if key=='notion_data_source_id' else 'google'
     connections.write(provider+':health',{'verified':False,'error':'Source changed. Refresh to verify access.'})
   log(None,'Source settings','Updated selected source IDs; authorisations unchanged')
   return {'saved':True}
 return await run_in_threadpool(perform)

@app.post('/api/connections/{provider}/start')
def start_connection(provider:str):
 try:url,cookie=connections.begin(provider)
 except ConnectionError as e:raise HTTPException(503,str(e))
 response=JSONResponse({'authorization_url':url})
 response.set_cookie('oauth_'+provider,cookie,max_age=600,httponly=True,secure=connections.base_url().startswith('https:'),samesite='lax',path='/auth/'+provider+'/callback')
 return response

@app.get('/auth/{provider}/callback')
def oauth_callback(provider:str,request:Request,code:str='',state:str='',error:str=''):
 try:
  with operation_lock:connections.finish(provider,state,request.cookies.get('oauth_'+provider,''),code,error)
  log(None,'Connection authorised',('Notion' if provider=='notion' else 'Google Drive')+' OAuth grant stored encrypted; awaiting source verification')
  location='/?connected='+provider
 except ConnectionError as e:
  # Error text is bounded and sanitised by the connection layer; no OAuth code is logged.
  location='/?connection_error='+urlencode({'message':str(e)})[8:]
 response=RedirectResponse(location,status_code=303)
 response.headers['Referrer-Policy']='no-referrer'
 response.headers['Cache-Control']='no-store'
 response.delete_cookie('oauth_'+provider,path='/auth/'+provider+'/callback')
 return response

@app.post('/api/connections/{provider}/disconnect')
def disconnect_connection(provider:str):
 if provider not in ('notion','google'):raise HTTPException(404,'Unknown provider')
 with operation_lock:connections.disconnect(provider)
 log(None,'Disconnected',provider+' portal access removed; revoke provider access separately if desired')
 return {'disconnected':True}

@app.post('/api/partners/publish-pending')
def publish_pending():
 with operation_lock:return {'results':services.publish_pending()}
