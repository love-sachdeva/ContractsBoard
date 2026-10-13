"""Deterministic document excerpts for human contract review."""
import io
import re
import uuid

from docx import Document
from pypdf import PdfReader


RULES = [
    ("Commission rates", r"commission|revenue share|reseller discount|margin", "Compare commission tiers and revenue definitions against the approved template."),
    ("Post-termination tail", r"post.termination|following termination|upon termination|commission payable", "Compare post-termination commission protection against the approved template."),
    ("Termination for convenience", r"termination for convenience|days.notice", "Compare notice periods and termination rights against the approved template."),
    ("Territory", r"territory|jurisdiction", "Compare the covered territory against the approved template."),
    ("Customer payments", r"non.payment|unable to recover|payment terms|payments from customer", "Compare payment timing and non-payment protection against the approved template."),
    ("Product change notice", r"discontinuation|product change|notice of changes", "Compare product-change notice requirements against the approved template."),
    ("Opportunity registration", r"approve or reject|opportunity registration|engagement request", "Compare opportunity registration rights against the approved template."),
    ("Renewals and upsells", r"renewals|upsells|extensions", "Compare renewal and expansion rights against the approved template."),
    ("Indemnification", r"indemni", "Compare scope, exclusions, and caps against the approved template."),
    ("Liability", r"limitation of liability|liability cap", "Compare caps, exclusions, and mutuality against the approved template."),
]

def extract(raw, name):
    if name.lower().endswith(".pdf"):
        parts = [(n + 1, page.extract_text() or "") for n, page in enumerate(PdfReader(io.BytesIO(raw)).pages)]
    elif name.lower().endswith(".docx"):
        document = Document(io.BytesIO(raw))
        text = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                text.append(" | ".join(cell.text for cell in row.cells))
        parts = [(1, "\n".join(text))]
    else:
        raise ValueError("Unsupported document type")
    if not any(text.strip() for _, text in parts):
        raise ValueError("No machine-readable text found. OCR/manual review required.")
    return parts

def snippets(parts, regex):
    pattern = re.compile(regex, re.I)
    found = []
    for page, text in parts:
        for match in list(pattern.finditer(text))[:3]:
            found.append(f"Page {page}: " + re.sub(r"\s+", " ", text[max(0, match.start() - 220): match.end() + 350]).strip()[:1100])
        if len(found) >= 3:
            break
    return "\n\n".join(found[:2])

def analyze_document(raw, name, doctype, services, config):
    if doctype not in ("LOI", "Reseller agreement"):
        return {"manual_review": True, "basis": "No baseline configured for this document type", "deviations": []}
    parts = extract(raw, name)
    template_id = config("LOI_TEMPLATE_ID" if doctype == "LOI" else "RESELLER_TEMPLATE_ID")
    base = None
    basis = "manual review checklist (approved template unavailable)"
    if template_id and services.connections.connected("google"):
        try:
            base = extract(services.file_bytes(template_id), "base.docx")
            basis = "live Drive template"
        except Exception:
            pass
    findings = []
    for subject, regex, checklist in RULES:
        partner_text = snippets(parts, regex)
        if not partner_text:
            continue
        base_text = snippets(base, regex) if base else checklist
        aligned = base and re.sub(r"\W+", "", partner_text.lower()) == re.sub(r"\W+", "", base_text.lower())
        findings.append({"id": str(uuid.uuid4()), "subject": subject, "base_text": base_text[:3500], "partner_text": partner_text[:3500], "decision": "Aligned" if aligned else "Needs review", "explanation": "Actual source extracts shown. Semantic differences require human assessment. Baseline: " + basis, "reviewer": "", "reviewed": None})
    return {"sections_extracted": len(parts), "review_items": len(findings), "basis": basis, "disclaimer": "Heuristic extraction, not legal review", "deviations": findings}
